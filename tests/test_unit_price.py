from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from sales_report.unit_price import _infer_latest_period, build_workbook


REFERENCE = Path("/workspace/scratch/7da2b172117b/reference_workbooks/사업계획 매출실적 분석/제품판매단가_원인탐색용_25년·26년_1~7월_통합.xlsx")


class UnitPriceWorkbookTests(unittest.TestCase):
    def test_eighth_month_expands_unit_price_scans(self):
        if not REFERENCE.exists():
            self.skipTest("reference template is only available in the work environment")
        headers = [
            "부문", "내수/수출", "요청월", "레벨1명", "레벨2명", "레벨3명", "계정", "중량",
            "달러금액", "한국원화금액", "담당자명", "담당자(세부)명", "약어명", "품명", "품번2",
        ]
        with TemporaryDirectory() as temp:
            root = Path(temp)
            raw_paths = []
            for year in (2025, 2026):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.append(headers)
                for month in range(1, 9):
                    worksheet.append(["합섬", "수출", f"{year}{month:02d}", "합섬방사", "", "", 2, 1000, 2000, 0, "북미", "북미", "TEST", "TEST", ""])
                path = root / f"{year}8월누계.xlsx"
                workbook.save(path)
                raw_paths.append(path)
            self.assertEqual(_infer_latest_period(raw_paths), (2026, 8))
            output = root / "result.xlsx"
            build_workbook(REFERENCE, raw_paths, output, 2026, 8)
            result = load_workbook(output, read_only=True, data_only=False)
            self.assertEqual(result["가공데이터"].max_row, 19)
            self.assertEqual(result["전년동월스캔"].cell(4, 32).value, "탐색키(최대변동)")
            self.assertEqual(result["변동스캔"].cell(4, 24).value, "탐색키(최대변동)")
            self.assertEqual(result["상세월별추이(선택)"].cell(3, 44).value, "8월 비중")
            self.assertIn("1~8월", result["사용가이드"]["A1"].value)

    def test_missing_month_does_not_stop_generation(self):
        if not REFERENCE.exists():
            self.skipTest("reference template is only available in the work environment")
        headers = [
            "부문", "내수/수출", "요청월", "레벨1명", "레벨2명", "레벨3명", "계정", "중량",
            "달러금액", "한국원화금액", "담당자명", "담당자(세부)명", "약어명",
        ]
        with TemporaryDirectory() as temp:
            root = Path(temp)
            raw_paths = []
            for year in (2025, 2026):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.append(headers)
                # 2026년 7월은 의도적으로 누락: 전년동월/전월비는 빈 값으로 남아야 함.
                for month in range(1, 9):
                    if year == 2026 and month == 7:
                        continue
                    worksheet.append(["합섬", "수출", f"{year}{month:02d}", "합섬방사", "", "", 2, 1000, 2000, 0, "북미", "북미", "TEST"])
                path = root / f"{year}8월누계.xlsx"
                workbook.save(path)
                raw_paths.append(path)
            output = root / "result.xlsx"
            build_workbook(REFERENCE, raw_paths, output, 2026, 8)
            self.assertTrue(output.exists())
